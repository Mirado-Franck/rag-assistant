"""API FastAPI du projet de classification et de recherche documentaire."""

from __future__ import annotations

import json
import re
import time
import unicodedata
import uuid
from collections import Counter
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

import logging
from fastapi import FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pydantic import BaseModel, Field

# --- INTÉGRATION JOUR 5 A1.2 : VRAI CLASSIFIEUR ---
# On garde le mock en fallback si Ollama est coupé
USE_REAL_CLASSIFIER = False
try:
    from src.rag.classifier import classify_with_alternatives
    USE_REAL_CLASSIFIER = True
    print("✅ Vrai classifieur chargé (src.rag.classifier)")
except Exception as e:
    print(f"⚠️ Classifieur réel indisponible ({e}) -> fallback mock")
    USE_REAL_CLASSIFIER = False

PROJECT_ROOT = Path(__file__).resolve().parents[2]
UPLOAD_DIR = PROJECT_ROOT / "data" / "uploads"
REGISTRY_FILE = PROJECT_ROOT / "data" / "documents_registry.json"
CHUNKS_FILE = PROJECT_ROOT / "data" / "chunks" / "all_chunks.json"
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
LLM_MODEL = "llama3.2"
EMBEDDING_MODEL = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"

CATEGORIES = [
    "Règlements",
    "Guides pratiques",
    "Procédures",
    "Organisation & plannings",
    "Formulaires & documents ponctuels",
]
ALLOWED_EXTENSIONS = {".pdf", ".txt", ".docx"}
# --- JOUR 6 A1 : LOGS ---
LOG_DIR = PROJECT_ROOT / "logs"
LOG_DIR.mkdir(parents=True, exist_ok=True)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(LOG_DIR / "app.log", encoding="utf-8"),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger("api")
logger.info("=== API DocuClassif AI démarrée ===")

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(
    title="DocuClassif AI — API",
    description="API de classification et de recherche dans les documents institutionnels",
    version="1.0.0",
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ---------------------------------------------------------------------------
# Modèles de requête/réponse
# ---------------------------------------------------------------------------
class ClassificationResult(BaseModel):
    categorie: str
    confiance: float
    justification: str
    alternatives: list[dict[str, Any]] = Field(default_factory=list)
    needs_confirmation: bool
    file_id: str
    filename: str
    text_preview: str


class ConfirmRequest(BaseModel):
    file_id: str
    categorie: str


class QueryRequest(BaseModel):
    question: str = Field(min_length=1, max_length=2000)
    category: str | None = None
    source_file: str | None = None
    n_results: int = Field(default=4, ge=1, le=20)


class QueryResponse(BaseModel):
    question: str
    reponse: str
    sources: list[dict[str, str]]
    filtre: str
    temps_ms: float
    nb_chunks_trouves: int


class DocumentInfo(BaseModel):
    filename: str
    categorie: str
    date_ajout: str
    type_original: str
    chunks_count: int


class HealthResponse(BaseModel):
    status: str
    modele_llm: str
    modele_embedding: str
    nb_documents: int
    nb_chunks: int
    categories: dict[str, int]


# ---------------------------------------------------------------------------
# Fichiers et données
# ---------------------------------------------------------------------------
def load_registry() -> dict[str, dict[str, Any]]:
    """Charge le registre, ou retourne un registre vide s'il n'existe pas."""
    if not REGISTRY_FILE.exists():
        return {}
    try:
        registry = json.loads(REGISTRY_FILE.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=500, detail="Le registre JSON est invalide.") from exc
    if not isinstance(registry, dict):
        raise HTTPException(status_code=500, detail="Le registre doit être un objet JSON.")
    return registry


def save_registry(registry: dict[str, dict[str, Any]]) -> None:
    """Écrit le registre de façon atomique pour éviter un JSON partiel."""
    REGISTRY_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = REGISTRY_FILE.with_suffix(".json.tmp")
    temporary_path.write_text(
        json.dumps(registry, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    temporary_path.replace(REGISTRY_FILE)


def load_chunks() -> list[dict[str, Any]]:
    """Charge les chunks disponibles pour les statistiques et la réponse mock."""
    if not CHUNKS_FILE.exists():
        return []
    try:
        chunks = json.loads(CHUNKS_FILE.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=500, detail="Le fichier de chunks JSON est invalide.") from exc
    if not isinstance(chunks, list):
        raise HTTPException(status_code=500, detail="Le fichier de chunks doit contenir une liste.")
    return chunks


def extract_uploaded_text(filename: str, content: bytes) -> tuple[str, int]:
    """Extrait le texte d'un PDF, DOCX ou TXT téléversé."""
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        try:
            import pymupdf

            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                page_texts = [page.get_text().strip() for page in pdf]
                return "\n\n".join(text for text in page_texts if text), len(page_texts)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"PDF illisible : {exc}") from exc

    if extension == ".docx":
        try:
            from docx import Document

            document = Document(BytesIO(content))
            blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        blocks.append(" | ".join(cells))
            # DOCX n'a pas de pagination stable sans moteur de mise en page.
            return "\n\n".join(blocks), 0
        except ImportError as exc:
            raise HTTPException(
                status_code=500,
                detail="Le support DOCX nécessite python-docx (pip install python-docx).",
            ) from exc
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"DOCX illisible : {exc}") from exc

    if extension == ".txt":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        return text, 1

    raise HTTPException(status_code=400, detail="Format non supporté.")


def classify_text_mock(text: str) -> dict[str, Any]:
    """Classification heuristique provisoire, en attendant l'intégration A1.

    Le titre du document est examiné avant le corps du texte : un règlement
    peut mentionner des dates/calendriers sans devenir un document de planning.
    """
    title_rules = [
        ("Règlements", ("règlement", "reglement", "code général", "code general", "statut")),
        ("Formulaires & documents ponctuels", ("formulaire", "attestation", "convention")),
        ("Procédures", ("procédure", "procedure", "démarche", "demarche")),
        ("Organisation & plannings", ("calendrier", "planning", "emploi du temps", "organigramme")),
        ("Guides pratiques", ("guide", "foire aux questions", "faq")),
    ]
    for line in (line.strip().casefold() for line in text.splitlines()[:15]):
        if not line or len(line) > 140:
            continue
        for category, keywords in title_rules:
            if any(keyword in line for keyword in keywords):
                return {
                    "categorie": category,
                    "confiance": 0.93,
                    "justification": "La catégorie est déterminée à partir du titre du document (classification simulée).",
                    "alternatives": [],
                    "needs_confirmation": False,
                }

    sample = text[:5000].casefold()
    rules = [
        (
            "Formulaires & documents ponctuels",
            ("formulaire", "attestation", "convention de stage", "nom :", "signature :"),
            0.91,
            "Le texte ressemble à un formulaire ou à un document à compléter.",
        ),
        (
            "Organisation & plannings",
            ("calendrier", "planning", "emploi du temps", "organigramme", "dates officielles"),
            0.87,
            "Le texte contient des dates ou des éléments d'organisation.",
        ),
        (
            "Règlements",
            ("article", "décret", "règlement", "interdit", "sanction", "obligation"),
            0.86,
            "Le texte contient des règles ou des formulations normatives.",
        ),
        (
            "Procédures",
            ("étape 1", "procédure", "démarche", "pièces à fournir", "à déposer"),
            0.82,
            "Le texte semble décrire une démarche à effectuer.",
        ),
        (
            "Guides pratiques",
            ("guide", "conseils", "comment", "foire aux questions", "faq"),
            0.76,
            "Le texte semble être un guide pratique.",
        ),
    ]

    for category, keywords, confidence, justification in rules:
        if any(keyword in sample for keyword in keywords):
            alternatives = []
            if confidence < 0.8:
                alternatives = [
                    {"categorie": "Procédures", "probabilite": 0.15},
                    {"categorie": "Règlements", "probabilite": 0.10},
                ]
            return {
                "categorie": category,
                "confiance": confidence,
                "justification": justification,
                "alternatives": alternatives,
                "needs_confirmation": confidence < 0.8,
            }

    return {
        "categorie": "Guides pratiques",
        "confiance": 0.45,
        "justification": "Classification simulée incertaine : une vérification humaine est nécessaire.",
        "alternatives": [
            {"categorie": "Procédures", "probabilite": 0.25},
            {"categorie": "Règlements", "probabilite": 0.15},
            {"categorie": "Formulaires & documents ponctuels", "probabilite": 0.10},
        ],
        "needs_confirmation": True,
    }


def classify_text_real(text: str) -> dict[str, Any]:
    """
    Classification RÉELLE via le LLM local (Ollama) - Jour 5 A1.2
    Remplace le mock du Jour 4. Garde le mock en fallback si Ollama est coupé.
    """
    if not USE_REAL_CLASSIFIER:
        return classify_text_mock(text)

    try:
        result = classify_with_alternatives(text, model_name="llama3.2")
        primary = result["primary"]
        alternatives = result.get("alternatives", [])
        needs_confirmation = result.get("needs_confirmation", False)

        return {
            "categorie": primary["categorie"],
            "confiance": primary["confiance"],
            "justification": primary["justification"],
            "alternatives": alternatives,
            "needs_confirmation": needs_confirmation,
        }

    except Exception as e:
        # Fallback : si le LLM ne répond pas, on ne fait pas crasher l'API
        print(f"⚠️ Erreur LLM classification: {e} -> fallback mock")
        try:
            # On tente le mock pour ne pas bloquer l'utilisateur
            mock_result = classify_text_mock(text)
            mock_result["justification"] = f"[Fallback mock - LLM erreur: {e}] {mock_result['justification']}"
            return mock_result
        except:
            return {
                "categorie": "Inconnu",
                "confiance": 0.0,
                "justification": f"Erreur de classification : {str(e)}",
                "alternatives": [],
                "needs_confirmation": True,
            }


def _normalize_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value.casefold())
    return "".join(char for char in normalized if not unicodedata.combining(char))


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------
@app.get("/api/health", response_model=HealthResponse, tags=["Système"])
def health_check() -> HealthResponse:
    """Retourne l'état de l'API et des données préparées."""
    registry = load_registry()
    chunks = load_chunks()
    counts = Counter(item.get("categorie", "Non classifié") for item in registry.values())
    return HealthResponse(
        status="ok",
        modele_llm=LLM_MODEL,
        modele_embedding=EMBEDDING_MODEL.rsplit("/", 1)[-1],
        nb_documents=len(registry),
        nb_chunks=len(chunks),
        categories={category: counts.get(category, 0) for category in CATEGORIES},
    )


@app.post("/api/upload", response_model=ClassificationResult, tags=["Documents"])
async def upload_document(file: UploadFile = File(...)) -> ClassificationResult:
    """Sauvegarde, extrait et classe avec le VRAI LLM (Jour 5 A1.2)."""
    if not file.filename:
        raise HTTPException(status_code=400, detail="Le fichier n'a pas de nom.")

    safe_name = Path(file.filename.replace(chr(92), "/")).name
    extension = Path(safe_name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Formats acceptés : PDF, TXT et DOCX.",
        )

    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Le fichier dépasse la limite de 20 Mo.")
    if not content:
        raise HTTPException(status_code=400, detail="Le fichier est vide.")

    try:
        text, _page_count = extract_uploaded_text(safe_name, content)
    finally:
        await file.close()

    if len(text.strip()) < 50:
        raise HTTPException(
            status_code=422,
            detail="Le fichier ne contient pas assez de texte exploitable (50 caractères minimum).",
        )

    file_id = uuid.uuid4().hex[:8]
    stored_name = f"{file_id}_{safe_name}"
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    (UPLOAD_DIR / stored_name).write_bytes(content)

    # --- JOUR 5 A1.2 : VRAI CLASSIFIEUR ---
    result = classify_text_real(text)

    return ClassificationResult(
        **result,
        file_id=file_id,
        filename=safe_name,
        text_preview=text[:300] + ("..." if len(text) > 300 else ""),
    )


@app.post("/api/confirm", tags=["Documents"])
def confirm_classification(request: ConfirmRequest) -> dict[str, Any]:
    """Confirme la catégorie et ajoute le document au registre partagé."""
    if request.categorie not in CATEGORIES:
        raise HTTPException(status_code=400, detail=f"Catégorie invalide. Choix : {CATEGORIES}")
    if not re.fullmatch(r"[0-9a-fA-F]{8}", request.file_id):
        raise HTTPException(status_code=400, detail="file_id invalide.")

    matches = list(UPLOAD_DIR.glob(f"{request.file_id}_*"))
    if len(matches) != 1 or not matches[0].is_file():
        raise HTTPException(status_code=404, detail="Fichier temporaire introuvable ou expiré.")

    stored_path = matches[0]
    original_name = stored_path.name[len(request.file_id) + 1 :]
    try:
        _, page_count = extract_uploaded_text(original_name, stored_path.read_bytes())
    except HTTPException:
        page_count = 0

    registry = load_registry()
    registry_key = stored_path.name
    registry_entry: dict[str, Any] = {
        "categorie": request.categorie,
        "date_ajout": date.today().isoformat(),
        "type_original": stored_path.suffix.lstrip(".").upper(),
        "chunks_count": 0,
        "status": "confirmé, indexation à effectuer",
    }
    # DOCX n'a pas de pagination fixe, et 0 pages invaliderait le registre.
    if page_count > 0:
        registry_entry["pages"] = page_count
    registry[registry_key] = registry_entry
    save_registry(registry)

    return {
        "status": "success",
        "message": (
            f"Document '{original_name}' confirmé et ajouté au registre. "
            "L'indexation ChromaDB sera connectée lors de l'intégration suivante (A2)."
        ),
        "file_id": request.file_id,
        "categorie": request.categorie,
        "indexe": False,
    }


@app.post("/api/query", response_model=QueryResponse, tags=["Recherche"])
def query_documents(request: QueryRequest) -> QueryResponse:
    """Réponse mock et extraits filtrés ; le vrai pipeline RAG est intégré plus tard (A2 Jour 5)."""
    started = time.perf_counter()
    if request.category and request.category not in CATEGORIES:
        raise HTTPException(status_code=400, detail=f"Catégorie invalide. Choix : {CATEGORIES}")

    registry = load_registry()
    eligible_chunks = []
    for chunk in load_chunks():
        source = chunk.get("source", "")
        metadata = registry.get(source, {})
        if request.category and metadata.get("categorie") != request.category:
            continue
        if request.source_file and source != request.source_file:
            continue
        eligible_chunks.append(chunk)

    query_terms = {
        token
        for token in re.findall(r"[\wÀ-ÿ]+", _normalize_text(request.question))
        if len(token) > 2
    }

    def score(chunk: dict[str, Any]) -> int:
        text = _normalize_text(str(chunk.get("text", "")))
        return sum(term in text for term in query_terms)

    ranked = sorted(eligible_chunks, key=score, reverse=True)[: request.n_results]
    sources: list[dict[str, str]] = []
    seen: set[tuple[str, int]] = set()
    for chunk in ranked:
        source = str(chunk.get("source", "inconnu"))
        chunk_index = int(chunk.get("chunk_index", 0))
        key = (source, chunk_index)
        if key in seen:
            continue
        seen.add(key)
        sources.append(
            {
                "fichier": source,
                "categorie": str(registry.get(source, {}).get("categorie", "Non classifié")),
                "extrait": str(chunk.get("text", ""))[:250],
            }
        )

    if sources:
        answer = (
            f"[MOCK] {len(sources)} extrait(s) trouvé(s). "
            "La génération de réponse par le pipeline RAG sera connectée lors de l'intégration (A2)."
        )
    else:
        answer = "Je ne trouve pas cette information dans les documents disponibles."

    if request.category and request.source_file:
        filter_label = f"catégorie={request.category} + fichier={request.source_file}"
    elif request.category:
        filter_label = f"catégorie={request.category}"
    elif request.source_file:
        filter_label = f"fichier={request.source_file}"
    else:
        filter_label = "aucun (tous les documents)"

    elapsed_ms = (time.perf_counter() - started) * 1000
    return QueryResponse(
        question=request.question,
        reponse=answer,
        sources=sources,
        filtre=filter_label,
        temps_ms=round(elapsed_ms, 1),
        nb_chunks_trouves=len(ranked),
    )


@app.get("/api/documents", response_model=list[DocumentInfo], tags=["Documents"])
def list_documents(category: str | None = Query(default=None)) -> list[DocumentInfo]:
    """Liste les documents du registre, avec filtre optionnel par catégorie."""
    if category and category not in CATEGORIES:
        raise HTTPException(status_code=400, detail=f"Catégorie invalide. Choix : {CATEGORIES}")

    registry = load_registry()
    chunk_counts = Counter(chunk.get("source", "") for chunk in load_chunks())
    documents = []
    for filename, info in sorted(registry.items()):
        if category and info.get("categorie") != category:
            continue
        documents.append(
            DocumentInfo(
                filename=filename,
                categorie=info.get("categorie", "Non classifié"),
                date_ajout=info.get("date_ajout", "inconnu"),
                type_original=info.get("type_original", "inconnu"),
                chunks_count=info.get("chunks_count", chunk_counts.get(filename, 0)),
            )
        )
    return documents


@app.get("/api/categories", tags=["Documents"])
def get_categories() -> dict[str, Any]:
    """Retourne les documents et le nombre de documents de chaque catégorie."""
    registry = load_registry()
    categories = {}
    for category in CATEGORIES:
        filenames = sorted(
            filename
            for filename, info in registry.items()
            if info.get("categorie") == category
        )
        categories[category] = {"nb_documents": len(filenames), "documents": filenames}
    return {"categories": categories, "total_documents": len(registry)}


# --- JOUR 6 A1 : GESTION D'ERREURS GLOBALE ---
@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    if isinstance(exc, HTTPException):
        # Laisser FastAPI gérer les HTTPException normalement
        raise exc
    logger.error(f"Erreur non gérée sur {request.url}: {exc}", exc_info=True)
    return JSONResponse(
        status_code=500,
        content={"erreur": "Une erreur interne est survenue", "detail": str(exc)}
    )

@app.exception_handler(HTTPException)
async def http_exception_handler(request: Request, exc: HTTPException):
    logger.warning(f"HTTP {exc.status_code} sur {request.url}: {exc.detail}")
    return JSONResponse(
        status_code=exc.status_code,
        content={"erreur": exc.detail, "status_code": exc.status_code}
    )

# Routes racine et santé historiques conservées pour la compatibilité.
@app.get("/health", include_in_schema=False)
def legacy_health() -> dict[str, str]:
    return {"status": "ok", "message": "DocuClassif AI - API fonctionnelle"}


@app.get("/", include_in_schema=False)
def root() -> dict[str, str]:
    return {"application": "DocuClassif AI", "version": "1.0.0", "documentation": "/docs"}
